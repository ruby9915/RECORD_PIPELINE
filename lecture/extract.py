"""강의자료·필기 텍스트 추출 (pdf, pptx, docx, hwpx, txt/md)"""
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .common import read_text_any

SUPPORTED = {".pdf": "pdf", ".pptx": "pptx", ".docx": "docx", ".hwpx": "hwpx", ".txt": "txt", ".md": "txt"}
UNSUPPORTED = {
    ".hwp": "구형 HWP(바이너리) 형식 — 한글에서 HWPX로 다시 저장하면 처리 가능",
    ".ppt": "구형 PPT 형식 — PowerPoint에서 PPTX로 저장하면 처리 가능",
    ".doc": "구형 DOC 형식 — Word에서 DOCX로 저장하면 처리 가능",
}
SEGMENT_CHARS = 2500  # 쪽 개념이 없는 문서는 이 길이 단위 '구간'으로 나눠 인용 위치를 만든다


def extract_document(path: Path) -> dict:
    """{"형식", "쪽단위": 쪽|슬라이드|구간, "쪽": [{"n": 1, "text": ..}], "경고": [..]}"""
    fmt = SUPPORTED.get(path.suffix.lower())
    if fmt == "pdf":
        pages, unit = _pdf(path), "쪽"
    elif fmt == "pptx":
        pages, unit = _pptx(path), "슬라이드"
    elif fmt == "docx":
        pages, unit = _segments(_docx(path)), "구간"
    elif fmt == "hwpx":
        pages, unit = _segments(_hwpx(path)), "구간"
    elif fmt == "txt":
        pages, unit = _segments(read_text_any(path).replace("\r", "").split("\n")), "구간"
    else:
        raise ValueError(f"지원하지 않는 형식: {path.suffix}")
    warnings = []
    total = sum(len(p["text"].strip()) for p in pages)
    if total == 0:
        warnings.append("추출된 텍스트 없음 (스캔 이미지 PDF이거나 빈 문서)")
    elif fmt in ("pdf", "pptx") and total < 40 * len(pages):
        warnings.append(f"텍스트가 매우 적음 ({total}자/{len(pages)}{unit}) — 이미지 위주 자료일 수 있음")
    return {"형식": fmt, "쪽단위": unit, "쪽": pages, "경고": warnings}


def _pdf(path):
    import pymupdf

    pages = []
    with pymupdf.open(str(path)) as doc:
        for i, page in enumerate(doc, 1):
            pages.append({"n": i, "text": _clean(page.get_text("text"))})
    return pages


def _pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_texts(shape):
        out = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                out += shape_texts(s)
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                out.append(t)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
        return out

    pages = []
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            parts += shape_texts(shape)
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame
            if nt is not None and nt.text.strip():
                parts.append("[발표자 노트] " + nt.text.strip())
        pages.append({"n": i, "text": _clean("\n".join(parts))})
    return pages


def _docx(path):
    import docx

    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return lines


def _hwpx(path):
    """HWPX = zip 안의 XML. 문단(hp:p) 단위로 텍스트(hp:t) 수집"""
    lines = []
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
        for n in names:
            root = ET.fromstring(z.read(n))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                    if text.strip():
                        lines.append(text)
    return lines


def _segments(lines):
    """쪽 개념이 없는 문서를 SEGMENT_CHARS 단위 '구간'으로 나눔"""
    pages, cur = [], ""
    for ln in lines:
        if cur and len(cur) + len(ln) + 1 > SEGMENT_CHARS:
            pages.append({"n": len(pages) + 1, "text": _clean(cur)})
            cur = ""
        cur = f"{cur}\n{ln}" if cur else ln
    if cur.strip() or not pages:
        pages.append({"n": len(pages) + 1, "text": _clean(cur)})
    return pages


def _clean(text):
    text = text.replace("\r", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def pages_text(doc: dict, pages=None) -> str:
    """자료 본문을 '=== p.N ===' 마커로 이어 붙임. pages: 포함할 쪽 번호 집합(None 이면 전체)"""
    out = []
    for p in doc["쪽"]:
        if pages and p["n"] not in pages:
            continue
        out.append(f"=== p.{p['n']} ===\n{p['text']}")
    return "\n".join(out)


def outline_text(doc: dict, head_chars=8000, per_page=150, total_cap=14000) -> str:
    """분류용 발췌: 앞부분은 전체, 나머지 쪽은 첫 줄 위주로"""
    out, used = [], 0
    for p in doc["쪽"]:
        t = p["text"]
        if used < head_chars:
            piece = t[: head_chars - used]
        else:
            piece = t[:per_page]
        out.append(f"=== p.{p['n']} ===\n{piece}")
        used += len(piece)
        if used >= total_cap:
            out.append(f"... (이하 {len(doc['쪽']) - p['n']}쪽 생략)")
            break
    return "\n".join(out)


def parse_page_range(spec, max_n):
    """'3-10, 15' → {3..10, 15}. '전체'/빈값 → None (전체)"""
    if spec is None or str(spec).strip() in ("", "전체", "all", "*"):
        return None
    s = set()
    for part in re.split(r"[,\s]+", str(spec).strip()):
        m = re.match(r"^p?\.?(\d+)\s*[-~–]\s*p?\.?(\d+)$", part)
        if m:
            a, b = int(m[1]), int(m[2])
            s.update(range(max(1, min(a, b)), min(max(a, b), max_n) + 1))
        else:
            m = re.match(r"^p?\.?(\d+)$", part)
            if m and 1 <= int(m[1]) <= max_n:
                s.add(int(m[1]))
    return s or None
